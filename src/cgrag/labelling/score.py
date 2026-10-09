"""Reads the filled-in workbooks back and turns them into numbers.

  job A  field-level precision / recall / F1 of the profile extractor, profile-level precision, agreement of the two labellers
  job B  a gold question file (intent, complexity, conditions, expected scope warning) for scripts/eval_questions.py
  job C  Cohen's kappa between the two labellers, the gold verdict per pair (agreement, else the third labeller), and how often the
         system's own verdict (private key file) matches the gold

Definitions (written down because the paper needs them):
  Job A, per field f, over profiles judged a real result (ALL_OK or SOME_WRONG):
    TP  the cell is filled and right                      FP  the cell is filled and wrong (flag WRONG)
    FN  the paper states a value that the profile does not have right: flag WRONG (a wrong value is also a miss) or flag MISSING
    precision = TP / (TP + FP)    recall = TP / (TP + FN)    F1 = their harmonic mean. Empty cells the paper does not state count for nothing.
  Rows judged NOT_A_RESULT are not part of the field scores; they lower the profile-level precision (real results / judged rows).
  CANNOT_CHECK rows are left out of everything.
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .sheets import A_FIELDS, A_FLAGS, A_VERDICTS, B_COMPLEXITY, B_CONDITIONS, B_INTENTS, C_DIFFERS, C_VERDICTS, a_flag_header

SYSTEM_VERDICT = {"GENUINE": "GENUINE", "EXPLAINED": "EXPLAINED", "NOT_COMPARABLE": "NOT COMPARABLE"}
SYSTEM_CONDITION = {"dataset_version": "dataset version / split", "model_size": "model size", "language": "language", "setting": "setting",
                    "task": "other"}
SPLIT_OR_SETTING = "split or setting"


def _clean(v) -> str | None:
    s = "" if v is None else str(v).strip()
    return s or None


_LIST_SPLIT = re.compile(r"\s*(?:;|,|\band\b|&)\s*", re.I)


def split_values(cell: str | None) -> list[str]:
    """A condition cell may name several things ("mBERT, XLM-R and GPT-4"): the list of them, in order, without repeats."""
    out: list[str] = []
    for part in _LIST_SPLIT.split(cell or ""):
        part = part.strip()
        if part and part.lower() not in {p.lower() for p in out}:
            out.append(part)
    return out


def read_rows(path: Path, sheet: str) -> list[dict]:
    """Rows of a sheet as dicts keyed by the header text (the header is row 1); completely empty rows are skipped."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(v is not None for v in r)]


def cohen_kappa(a: list[str], b: list[str]) -> float:
    """Cohen's kappa for two labellers over the same items (1.0 = identical, 0 = no better than chance)."""
    assert len(a) == len(b) and a, "two equally long, non-empty label lists are needed"
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / (n * n)
    return 1.0 if pe == 1.0 else (po - pe) / (1 - pe)


def prf(tp: int, fp: int, fn: int) -> dict:
    p = tp / (tp + fp) if tp + fp else None
    r = tp / (tp + fn) if tp + fn else None
    f = 2 * p * r / (p + r) if p and r else (0.0 if tp + fp + fn else None)
    return {"tp": tp, "fp": fp, "fn": fn, "precision": None if p is None else round(p, 4), "recall": None if r is None else round(r, 4),
            "f1": None if f is None else round(f, 4)}


# ---------------------------------------------------------------- job A ----

def parse_job_a(path: Path) -> tuple[list[dict], list[str]]:
    """(rows, problems). A row: profile_id, paper, verdict, flags {field: WRONG|MISSING|None}, fields {field: value|None}, note, file."""
    rows, problems = [], []
    for i, r in enumerate(read_rows(path, "Profiles"), start=2):
        verdict = _clean(r.get("verdict"))
        pid = _clean(r.get("profile_id")) or f"row{i}"
        if verdict is None:
            continue
        if verdict not in A_VERDICTS:
            problems.append(f"{path.name} row {i}: unknown verdict {verdict!r}")
            continue
        flags = {f: _clean(r.get(a_flag_header(f))) for f in A_FIELDS}
        for f, v in flags.items():
            if v is not None and v not in A_FLAGS:
                problems.append(f"{path.name} row {i}: {f} flag {v!r} is not WRONG or MISSING")
                flags[f] = None
        if verdict == "SOME_WRONG" and not any(flags.values()):
            problems.append(f"{path.name} row {i}: SOME_WRONG but no cell is marked WRONG or MISSING")
        if verdict in ("ALL_OK", "NOT_A_RESULT", "CANNOT_CHECK") and any(flags.values()):
            problems.append(f"{path.name} row {i}: {verdict} but a cell is marked - the marks are ignored")
        rows.append({"profile_id": pid, "paper": (_clean(r.get("paper")) or "").split()[0] if _clean(r.get("paper")) else "",
                     "verdict": verdict, "flags": flags, "fields": {f: _clean(r.get(f)) for f in A_FIELDS},
                     "note": _clean(r.get("note / correct value")), "file": path.name})
    return rows, problems


def score_job_a(rows: list[dict], group_of: dict[str, str] | None = None) -> dict:
    """Field-level and profile-level scores. A profile labelled by two people counts once (the first file's label)."""
    seen, unique = set(), []
    for r in rows:
        if r["profile_id"] not in seen:
            seen.add(r["profile_id"])
            unique.append(r)

    def block(part: list[dict]) -> dict:
        counts = Counter(r["verdict"] for r in part)
        judged = counts["ALL_OK"] + counts["SOME_WRONG"] + counts["NOT_A_RESULT"]
        tp, fp, fn = defaultdict(int), defaultdict(int), defaultdict(int)
        for r in part:
            if r["verdict"] not in ("ALL_OK", "SOME_WRONG"):
                continue
            for f in A_FIELDS:
                filled = r["fields"][f] is not None
                flag = r["flags"][f] if r["verdict"] == "SOME_WRONG" else None
                if flag == "WRONG" and filled:
                    fp[f] += 1
                    fn[f] += 1
                elif flag in ("WRONG", "MISSING"):                 # MISSING (or WRONG on an empty cell): the value was not extracted
                    fn[f] += 1
                elif filled:
                    tp[f] += 1
        per_field = {f: prf(tp[f], fp[f], fn[f]) for f in A_FIELDS}
        micro = prf(sum(tp.values()), sum(fp.values()), sum(fn.values()))
        return {"profiles_labelled": len(part), "verdicts": dict(counts),
                "profile_precision": round((counts["ALL_OK"] + counts["SOME_WRONG"]) / judged, 4) if judged else None,
                "fully_correct_of_results": round(counts["ALL_OK"] / (counts["ALL_OK"] + counts["SOME_WRONG"]), 4)
                if counts["ALL_OK"] + counts["SOME_WRONG"] else None,
                "fully_correct_of_judged": round(counts["ALL_OK"] / judged, 4) if judged else None,
                "per_field": per_field, "micro": micro}

    out = {"all": block(unique)}
    if group_of:
        for g in sorted(set(group_of.values())):
            out[g] = block([r for r in unique if group_of.get(r["paper"]) == g])
    return out


def agreement_job_a(rows: list[dict]) -> dict:
    """Agreement of two labellers on the profiles both labelled (the shared rows)."""
    by_profile: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_profile[r["profile_id"]].append(r)
    both = {pid: rs for pid, rs in by_profile.items() if len({r["file"] for r in rs}) >= 2}
    if not both:
        return {"shared_profiles": 0}
    first, second = [], []
    cells_equal = cells_total = 0
    for rs in both.values():
        a, b = rs[0], next(r for r in rs[1:] if r["file"] != rs[0]["file"])
        first.append(a["verdict"])
        second.append(b["verdict"])
        for f in A_FIELDS:
            cells_total += 1
            cells_equal += (a["flags"][f] if a["verdict"] == "SOME_WRONG" else None) == (b["flags"][f] if b["verdict"] == "SOME_WRONG" else None)
    return {"shared_profiles": len(both), "verdict_agreement": round(sum(x == y for x, y in zip(first, second)) / len(first), 4),
            "verdict_kappa": round(cohen_kappa(first, second), 4), "cell_mark_agreement": round(cells_equal / cells_total, 4)}


# ---------------------------------------------------------------- job B ----

def parse_job_b(path: Path) -> tuple[list[dict], list[str]]:
    """(questions, problems): one dict per written question, in the gold-file format."""
    out, problems = [], []
    for i, r in enumerate(read_rows(path, "Questions"), start=2):
        question = _clean(r.get("question"))
        if question is None:
            continue
        intent, cx, warn = _clean(r.get("intent")), _clean(r.get("complexity")), _clean(r.get("expected scope warning"))
        for label, value, allowed in (("intent", intent, B_INTENTS), ("complexity", cx, B_COMPLEXITY), ("expected scope warning", warn, ["YES", "NO"])):
            if value not in allowed:
                problems.append(f"{path.name} row {i} ({_clean(r.get('id'))}): {label} is {value!r}, expected one of {allowed}")
        conditions = {c: split_values(_clean(r.get(c))) for c in B_CONDITIONS if _clean(r.get(c))}
        missing = [m.strip() for m in (_clean(r.get("missing conditions")) or "").replace(";", ",").split(",") if m.strip()]
        if warn == "YES" and not missing:
            problems.append(f"{path.name} row {i} ({_clean(r.get('id'))}): a scope warning is expected but no missing condition is named")
        planned = _clean(r.get("planned type"))
        if planned == "covered" and warn == "YES":
            problems.append(f"{path.name} row {i} ({_clean(r.get('id'))}): planned 'covered' but a warning is expected - check the question")
        if planned in ("one missing", "partly covered") and warn == "NO":
            problems.append(f"{path.name} row {i} ({_clean(r.get('id'))}): planned {planned!r} but no warning is expected - check the question")
        out.append({"id": _clean(r.get("id")), "author": _clean(r.get("author")), "planned_type": planned, "question": question,
                    "intent": intent, "complexity": cx, "conditions": conditions, "expect_warning": warn == "YES", "missing": missing,
                    "evidence": _clean(r.get("evidence (paper, page, table)")), "facts": _clean(r.get("expected answer facts")),
                    "notes": _clean(r.get("notes"))})
    return out, problems


def write_jsonl(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), encoding="utf-8")


# ---------------------------------------------------------------- job C ----

def _verdict(v: str | None) -> str | None:
    return None if v is None else v.strip().upper().replace("_", " ")


def parse_job_c(path: Path) -> tuple[dict[str, dict], list[str]]:
    """({pair_id: label}, problems). label: verdict, differs (list of condition names), other, extraction_error, confidence, explanation."""
    out, problems = {}, []
    for i, r in enumerate(read_rows(path, "Pairs"), start=2):
        pid = _clean(r.get("pair"))
        verdict = _verdict(_clean(r.get("verdict")))
        if not pid or verdict is None:
            continue
        if verdict not in C_VERDICTS:
            problems.append(f"{path.name} row {i} ({pid}): unknown verdict {verdict!r}")
            continue
        differs = [d for d in C_DIFFERS if _clean(r.get(f"differs: {d}"))]
        if verdict == "EXPLAINED" and not differs:
            problems.append(f"{path.name} row {i} ({pid}): EXPLAINED but no condition is ticked")
        if verdict != "EXPLAINED" and differs:
            problems.append(f"{path.name} row {i} ({pid}): {verdict} but conditions are ticked - the ticks are ignored")
            differs = []
        out[pid] = {"verdict": verdict, "differs": differs, "other": _clean(r.get("other: what")),
                    "extraction_error": bool(_clean(r.get("extraction error?"))), "confidence": _clean(r.get("confidence")),
                    "explanation": _clean(r.get("explanation (quote the paper)"))}
    return out, problems


def agreement_job_c(first: dict[str, dict], second: dict[str, dict]) -> dict:
    ids = sorted(set(first) & set(second))
    if not ids:
        return {"pairs_labelled_by_both": 0}
    a, b = [first[i]["verdict"] for i in ids], [second[i]["verdict"] for i in ids]
    explained = [i for i in ids if first[i]["verdict"] == second[i]["verdict"] == "EXPLAINED"]
    same_conditions = sum(set(first[i]["differs"]) == set(second[i]["differs"]) for i in explained)
    kappa = cohen_kappa(a, b)
    return {"pairs_labelled_by_both": len(ids), "percent_agreement": round(sum(x == y for x, y in zip(a, b)) / len(ids), 4),
            "kappa": round(kappa, 4), "meets_target_0.6": kappa >= 0.6, "disagreements": [i for i in ids if first[i]["verdict"] != second[i]["verdict"]],
            "both_say_explained": len(explained), "same_conditions_ticked": same_conditions,
            "verdicts_first": dict(Counter(a)), "verdicts_second": dict(Counter(b))}


def gold_job_c(first: dict[str, dict], second: dict[str, dict], third: dict[str, dict] | None = None) -> dict[str, dict]:
    """The gold label of every pair both labellers did: their verdict when they agree, otherwise the third labeller's (a pair they
    disagree on and nobody settled is left out). For EXPLAINED the differing conditions are those every voter for that verdict ticked."""
    gold = {}
    for pid in sorted(set(first) & set(second)):
        votes = [first[pid], second[pid]] + ([third[pid]] if third and pid in third else [])
        if first[pid]["verdict"] == second[pid]["verdict"]:
            verdict = first[pid]["verdict"]
        elif third and pid in third:
            verdict = third[pid]["verdict"]
        else:
            continue
        voters = [v for v in votes if v["verdict"] == verdict]
        differs = [d for d in C_DIFFERS if all(d in v["differs"] for v in voters)] if verdict == "EXPLAINED" else []
        gold[pid] = {"verdict": verdict, "differs": differs, "extraction_error": any(v["extraction_error"] for v in votes),
                     "settled_by_third": first[pid]["verdict"] != second[pid]["verdict"]}
    return gold


def score_system_c(gold: dict[str, dict], key: dict[str, dict]) -> dict:
    """The system's verdict (private key file) against the gold: accuracy, per-class precision/recall/F1, confusion, condition attribution."""
    ids = [i for i in gold if i in key]
    classes = C_VERDICTS
    confusion = {g: {s: 0 for s in classes} for g in classes}
    for i in ids:
        confusion[gold[i]["verdict"]][SYSTEM_VERDICT[key[i]["system_verdict"]]] += 1
    per_class = {}
    for c in classes:
        tp = confusion[c][c]
        fp = sum(confusion[g][c] for g in classes if g != c)
        fn = sum(confusion[c][s] for s in classes if s != c)
        per_class[c] = {**prf(tp, fp, fn), "support": sum(confusion[c].values())}
    f1s = [v["f1"] for v in per_class.values() if v["f1"] is not None and v["support"]]
    attribution, attribution_merged = [], []
    # The labelling instructions list "dev vs test" under two names ("dataset version / split" and "setting"), so two labellers (or a labeller
    # and the system) can name the same difference differently without either being wrong. The merged reading treats the two names as one.
    merge = lambda names: {SPLIT_OR_SETTING if n in ("dataset version / split", "setting") else n for n in names}
    for i in ids:
        if gold[i]["verdict"] == "EXPLAINED" and SYSTEM_VERDICT[key[i]["system_verdict"]] == "EXPLAINED":
            sys_set = {SYSTEM_CONDITION.get(d, "other") for d in key[i]["system_differing"]}
            gold_set = set(gold[i]["differs"])
            attribution.append((sys_set == gold_set, len(sys_set & gold_set) / len(sys_set | gold_set) if sys_set | gold_set else 1.0))
            m_sys, m_gold = merge(sys_set), merge(gold_set)
            attribution_merged.append((m_sys == m_gold, len(m_sys & m_gold) / len(m_sys | m_gold) if m_sys | m_gold else 1.0))
    return {"pairs": len(ids), "accuracy": round(sum(confusion[c][c] for c in classes) / len(ids), 4) if ids else None,
            "macro_f1": round(sum(f1s) / len(f1s), 4) if f1s else None, "per_class": per_class, "confusion_gold_rows_system_columns": confusion,
            "condition_attribution": {"pairs": len(attribution), "exact_match": round(sum(a for a, _ in attribution) / len(attribution), 4) if attribution else None,
                                      "mean_jaccard": round(sum(j for _, j in attribution) / len(attribution), 4) if attribution else None},
            "condition_attribution_merged": {
                "pairs": len(attribution_merged), "note": "'dataset version / split' and 'setting' counted as one name (dev vs test is listed under both)",
                "exact_match": round(sum(a for a, _ in attribution_merged) / len(attribution_merged), 4) if attribution_merged else None,
                "mean_jaccard": round(sum(j for _, j in attribution_merged) / len(attribution_merged), 4) if attribution_merged else None},
            "extraction_errors_flagged": sum(gold[i]["extraction_error"] for i in ids)}
