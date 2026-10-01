"""openpyxl helpers shared by the three labelling workbooks: one look, dropdowns, locked cells, a READ ME sheet.

Colour legend used in every workbook: YELLOW = you type or pick here; GREY = given by the system, do not change; GREEN = row done.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=15, bold=True, color="1F3864")
F_H2 = Font(name=FONT, size=11, bold=True, color="1F3864")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2A8")        # yellow: type or pick here
FILL_FIXED = PatternFill("solid", fgColor="F2F2F2")        # grey: given, do not change
FILL_DONE = PatternFill("solid", fgColor="D9EAD3")         # green: this row is filled in
FILL_EXAMPLE = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)      # Excel computes the progress counters when the file opens
    return wb


def col(n: int) -> str:
    return get_column_letter(n)


def header_row(ws: Worksheet, headers: list[str], widths: list[float], row: int = 1) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[col(i)].width = w
    ws.row_dimensions[row].height = 42


def put(ws: Worksheet, row: int, column: int, value, *, kind: str = "fixed", link: str | None = None, center: bool = False,
        number_format: str | None = None):
    """kind: 'fixed' (grey, locked), 'input' (yellow, unlocked), 'plain' (no fill, locked)."""
    c = ws.cell(row=row, column=column, value=value)
    c.font = F_LINK if link else F_BASE
    c.alignment = CENTER if center else WRAP
    c.border = BORDER
    if link:
        c.hyperlink = link
    if kind == "input":
        c.fill, c.protection = FILL_INPUT, Protection(locked=False)
    elif kind == "fixed":
        c.fill = FILL_FIXED
    if number_format:
        c.number_format = number_format
    return c


def wrapped_lines(text: str, width: float) -> int:
    """How many lines a text takes in a column of this width (a rough count: ~1.15 characters per width unit)."""
    per_line = max(int(width * 1.15), 8)
    return sum(max(1, -(-len(part) // per_line)) for part in str(text).split("\n"))


def fit_row(ws: Worksheet, row: int, texts_and_widths: list[tuple[str, float]], *, minimum: float = 30, maximum: float = 260) -> None:
    lines = max((wrapped_lines(t, w) for t, w in texts_and_widths if t), default=1)
    ws.row_dimensions[row].height = max(minimum, min(maximum, 13.2 * lines + 6))


def dropdown(ws: Worksheet, cell_range: str, options: list[str], *, title: str, prompt: str) -> None:
    formula = '"' + ",".join(options) + '"'
    assert len(formula) <= 255, "an inline dropdown list is limited to 255 characters"
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True, showInputMessage=True,
                        errorTitle=title, error="Pick one of: " + ", ".join(options), promptTitle=title, prompt=prompt)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def highlight_done(ws: Worksheet, cell_range: str, first_cell: str) -> None:
    """Green fill on a cell of the range once it is not empty (first_cell is the top-left cell of the range, e.g. P2)."""
    ws.conditional_formatting.add(cell_range, FormulaRule(formula=[f'LEN({first_cell})>0'], fill=FILL_DONE))


def protect(ws: Worksheet, hidden: tuple[int, ...] = ()) -> None:
    """Lock everything except the unlocked (yellow) cells; filtering, resizing and hiding stay possible. No password."""
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False
    for n in hidden:
        ws.column_dimensions[col(n)].hidden = True


def readme_sheet(wb: Workbook, title: str, blocks: list[tuple[str, list[str]]], *, legend: bool = True,
                 example: tuple[list[str], list[list[str]]] | None = None, progress: str | None = None) -> Worksheet:
    """The first sheet: what this job is, what to do, definitions, a colour legend, an example, and a progress counter.

    blocks: (heading, lines) pairs. example: (headers, rows) shown as a small table. progress: a formula (without '=') for a live
    'rows labelled' counter."""
    ws = wb.active
    ws.title = "READ ME"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for letter, width in zip("BCDEFGHI", (30, 24, 24, 24, 24, 24, 24, 24)):
        ws.column_dimensions[letter].width = width
    row = 1
    ws.cell(row=row, column=2, value=title).font = F_TITLE
    row += 1
    if progress:
        c = ws.cell(row=row, column=2, value="=" + progress)
        c.font = F_BOLD
        row += 1
    row += 1
    for heading, lines in blocks:
        ws.cell(row=row, column=2, value=heading).font = F_H2
        row += 1
        for line in lines:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
            c = ws.cell(row=row, column=2, value=line)
            c.font, c.alignment = F_BASE, WRAP
            ws.row_dimensions[row].height = max(15, 13.2 * wrapped_lines(line, 175) + 3)
            row += 1
        row += 1
    if legend:
        ws.cell(row=row, column=2, value="Colours in the data sheet").font = F_H2
        row += 1
        for fill, text in ((FILL_INPUT, "YELLOW = you type or pick here (the only cells you change)"),
                           (FILL_FIXED, "GREY = given by the system - do not change"),
                           (FILL_DONE, "GREEN = this row is filled in")):
            c = ws.cell(row=row, column=2, value=text.split(" = ")[0])
            c.fill, c.font, c.border = fill, F_BOLD, BORDER
            d = ws.cell(row=row, column=3, value=text.split(" = ")[1])
            d.font = F_BASE
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
            row += 1
        row += 1
    if example:
        headers, rows = example
        ws.cell(row=row, column=2, value="Example of a finished row (this one is NOT in your data sheet)").font = F_H2
        row += 1
        for r in (headers, *rows):
            for i, v in enumerate(r):
                c = ws.cell(row=row, column=2 + i, value=v)
                c.font = F_BOLD if r is headers else F_BASE
                c.fill = FILL_HEAD if r is headers else FILL_EXAMPLE
                if r is headers:
                    c.font = F_HEAD
                c.alignment, c.border = WRAP, BORDER
            ws.row_dimensions[row].height = max(18, 13.2 * max(wrapped_lines(str(v), 24) for v in r) + 4)
            row += 1
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    return ws
