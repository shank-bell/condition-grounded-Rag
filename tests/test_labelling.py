"""The labelling kit: sampling, the workbooks, and a full round trip (build -> fill in like a labeller -> parse -> score)."""
import random

import pytest
from openpyxl import load_workbook

from cgrag.labelling.sampling import (Pair, corpus_map, evidence_excerpt, language_label, mine_pairs, pick_pairs, sample_profiles,
                                      value_forms)
from cgrag.labelling.score import (agreement_job_a, agreement_job_c, cohen_kappa, gold_job_c, parse_job_a, parse_job_b, parse_job_c,
                                   score_job_a, score_system_c)
from cgrag.labelling.sheets import (A_FIELDS, B_HEADERS, C_HEADERS, a_flag_header, a_rows, build_job_a, build_job_b, build_job_c)
from cgrag.schemas import ConditionProfile


def prof(i: int, paper="p1", chunk="p1:1", **kw) -> ConditionProfile:
    base = dict(profile_id=f"{chunk}#{i}", paper_id=paper, chunk_id=chunk, metric="accuracy", value=50.0 + i, model="mBERT", dataset="XNLI")
    return ConditionProfile(**{**base, **kw})


TABLE = """Table 3: XNLI results (accuracy).
| Model | en | hi | kn |
| mBERT | 81.4 | 60.1 | 58.6 |
| XLM-R | 85.8 | 71.9 | 71.5 |"""


# ---------- statistics and excerpts ----------

def test_cohen_kappa_matches_a_worked_example():
    a = ["Y"] * 20 + ["Y"] * 5 + ["N"] * 10 + ["N"] * 15
    b = ["Y"] * 20 + ["N"] * 5 + ["Y"] * 10 + ["N"] * 15
    assert cohen_kappa(a, b) == pytest.approx(0.4)
    assert cohen_kappa(a, a) == 1.0
    assert cohen_kappa(["a", "b"] * 5, ["b", "a"] * 5) == pytest.approx(-1.0)


def test_value_forms_cover_how_a_table_writes_a_number():
    assert set(value_forms(56.0)) >= {"56", "56.0", "56.00"}
    assert "83.10" in value_forms(83.1)


def test_the_excerpt_marks_the_number_and_picks_the_row_of_the_model():
    p = prof(1, value=71.5, model="XLM-R", language="Kannada")
    text = evidence_excerpt(TABLE, p)
    assert text.splitlines()[0].startswith("Table 3") and "| Model | en | hi | kn |" in text
    assert "XLM-R" in text and "[[71.5]]" in text and "mBERT" not in text.splitlines()[-1]


def test_a_wide_table_keeps_the_value_s_column_visible():
    header = "| Model | " + " | ".join(f"lang{i}" for i in range(40)) + " |"
    row = "| XLM-R | " + " | ".join(f"{60 + i * 0.1:.1f}" for i in range(40)) + " |"
    text = evidence_excerpt(f"Table 9: big\n{header}\n{row}", prof(1, value=62.5, model="XLM-R"), max_line=120)
    assert "[[62.5]]" in text and "lang25" in text and len(max(text.splitlines(), key=len)) < 200


def test_a_prose_chunk_gets_the_passage_around_the_number():
    text = evidence_excerpt("We find that BERT-large reaches 90.9 F1 on the SQuAD dev set.", prof(1, value=90.9))
    assert "[[90.9]]" in text and "SQuAD" in text
    assert "not found" in evidence_excerpt("nothing here", prof(1, value=90.9))


# ---------- sampling ----------

def test_profiles_are_sampled_per_paper_with_a_cap_per_chunk():
    profiles = [prof(i, paper="a", chunk="a:1") for i in range(30)] + [prof(i, paper="a", chunk="a:2") for i in range(30)] \
        + [prof(i, paper="b", chunk=f"b:{i % 5}") for i in range(30)]
    got = sample_profiles(profiles, ["a", "b"], 10, random.Random(1), cap_per_chunk=3)
    assert sum(p.paper_id == "a" for p in got) == 10 and sum(p.paper_id == "b" for p in got) == 10
    per_chunk = {c: sum(1 for p in got if p.chunk_id == c) for c in {p.chunk_id for p in got}}
    assert all(per_chunk[f"b:{k}"] <= 3 for k in range(5) if f"b:{k}" in per_chunk)                    # five chunks: the cap of 3 holds
    assert per_chunk["a:1"] >= 3 and per_chunk["a:2"] >= 3                                             # two chunks: both are seen; the cap is lifted to reach 10


def test_pairs_are_mined_across_papers_and_split_into_strata():
    a = prof(1, paper="a", chunk="a:1", model="BERT-large", dataset="SQuAD", dataset_version="1.1", metric="F1", value=90.9, setting="dev set")
    b = prof(2, paper="b", chunk="b:1", model="BERT-large", dataset="SQuAD", dataset_version="2.0", metric="F1", value=81.8, setting="dev set")
    c = prof(3, paper="c", chunk="c:1", model="BERT-large", dataset="SQuAD", dataset_version="1.1", metric="F1", value=90.9)   # same as a: no pair
    d = prof(4, paper="d", chunk="d:1", model="RoBERTa", dataset="SQuAD", metric="F1", value=50.0)                             # another model
    pairs = mine_pairs([a, b, c, d])
    assert {(p.x.paper_id, p.y.paper_id) for p in pairs} == {("a", "b"), ("b", "c")}
    assert {p.stratum for p in pairs} == {"explained_version"}
    assert len(pick_pairs(pairs, 5, random.Random(0))) == 2


def test_languages_are_named_and_junk_is_dropped():
    assert language_label("kn") == "Kannada" and language_label("hindi") == "Hindi"
    assert language_label("7 languages") is None and language_label("multi-language") is None and language_label(None) is None


def test_the_corpus_map_shows_which_language_each_dataset_has():
    profiles = [prof(i, dataset="IndicXNLI", language=lang, paper="x", chunk="x:1") for i in range(12) for lang in ("kn", "Hindi")] \
        + [prof(100 + i, dataset="XNLI", language="Hindi", paper="y", chunk="y:1") for i in range(20)] \
        + [prof(200 + i, dataset="XNLI", language="French", paper="y", chunk="y:1") for i in range(5)]
    sheets = corpus_map(profiles, {"x": "A paper", "y": "Another"})
    headers, rows = sheets["Dataset x Language"]
    by_dataset = {r[0]: dict(zip(headers, r)) for r in rows}
    assert by_dataset["IndicXNLI"]["Kannada"] == 12 and by_dataset["XNLI"]["Kannada"] is None and by_dataset["XNLI"]["French"] == 5
    assert set(sheets) == {"Papers", "Datasets", "Dataset x Language", "Models", "Model x Dataset"}
    assert sheets["Papers"][1][0][-1][1].startswith("https://arxiv.org/pdf/x#page=1")


# ---------- job A round trip ----------

def column_of(ws, header: str) -> int:
    return next(c.column for c in ws[1] if c.value == header)


def fill(path, sheet, row_values: dict[int, dict[str, object]]) -> None:
    wb = load_workbook(path)
    ws = wb[sheet]
    for row, values in row_values.items():
        for header, v in values.items():
            ws.cell(row=row, column=column_of(ws, header), value=v)
    wb.save(path)


def test_job_a_workbook_has_the_promised_structure(tmp_path):
    rows = a_rows([prof(1, value=71.5, model="XLM-R", chunk="p1:1")], {"p1:1": TABLE}, {"p1:1": 7}, {"p1": "Some title"})
    path = build_job_a(tmp_path / "a.xlsx", "Adarsh", rows)
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "READ ME" and "Profiles" in wb.sheetnames
    ws = wb["Profiles"]
    assert ws.protection.sheet and ws.freeze_panes and ws.column_dimensions["B"].hidden
    assert ws.cell(row=2, column=column_of(ws, "open")).hyperlink.target == "https://arxiv.org/pdf/p1#page=7"
    assert "[[71.5]]" in ws.cell(row=2, column=column_of(ws, "source (caption, header, the row the number came from)")).value
    verdict = ws.cell(row=2, column=column_of(ws, "verdict"))
    assert verdict.fill.fgColor.rgb.endswith("FFF2A8") and verdict.protection.locked is False       # yellow and editable
    assert ws.cell(row=2, column=column_of(ws, "model")).protection.locked is True                    # grey and locked
    assert any("ALL_OK" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    readme = wb["READ ME"]
    assert any("Rows done" in str(c.value) for r in readme.iter_rows() for c in r if c.value)


def test_job_a_round_trip_scores_fields(tmp_path):
    profiles = [prof(i, language="Hindi" if i < 4 else None, task="natural language inference", chunk="p1:1") for i in range(1, 6)]
    path = build_job_a(tmp_path / "a.xlsx", "Adarsh", a_rows(profiles, {"p1:1": TABLE}, {"p1:1": 3}, {}))
    fill(path, "Profiles", {
        2: {"verdict": "ALL_OK"},
        3: {"verdict": "SOME_WRONG", a_flag_header("language"): "WRONG"},
        4: {"verdict": "SOME_WRONG", a_flag_header("setting"): "MISSING"},
        5: {"verdict": "NOT_A_RESULT"},
        6: {"verdict": "CANNOT_CHECK"},
    })
    rows, problems = parse_job_a(path)
    assert problems == [] and len(rows) == 5
    score = score_job_a(rows)["all"]
    assert score["verdicts"] == {"ALL_OK": 1, "SOME_WRONG": 2, "NOT_A_RESULT": 1, "CANNOT_CHECK": 1}
    assert score["profile_precision"] == pytest.approx(3 / 4)                                   # 3 real results out of 4 judged
    lang, setting, value = score["per_field"]["language"], score["per_field"]["setting"], score["per_field"]["value"]
    assert (lang["tp"], lang["fp"], lang["fn"]) == (2, 1, 1)                                   # rows 2 and 4 right, row 3 wrong
    assert (setting["tp"], setting["fn"]) == (0, 1)                                            # MISSING counts as a miss only
    assert (value["tp"], value["fp"], value["fn"]) == (3, 0, 0)


def test_job_a_problems_are_reported(tmp_path):
    path = build_job_a(tmp_path / "a.xlsx", "Adarsh", a_rows([prof(1), prof(2)], {}, {}, {}))
    fill(path, "Profiles", {2: {"verdict": "SOME_WRONG"}, 3: {"verdict": "MAYBE"}})
    rows, problems = parse_job_a(path)
    assert len(rows) == 1 and any("no cell is marked" in p for p in problems) and any("unknown verdict" in p for p in problems)


def test_job_a_agreement_on_shared_profiles(tmp_path):
    profiles = [prof(i, chunk="p1:1") for i in range(1, 5)]
    for name, verdicts in (("x", ["ALL_OK", "ALL_OK", "SOME_WRONG", "ALL_OK"]), ("y", ["ALL_OK", "SOME_WRONG", "SOME_WRONG", "ALL_OK"])):
        path = build_job_a(tmp_path / f"{name}.xlsx", name, a_rows(profiles, {}, {}, {}))
        fill(path, "Profiles", {i + 2: {"verdict": v, **({a_flag_header("model"): "WRONG"} if v == "SOME_WRONG" else {})} for i, v in enumerate(verdicts)})
    rows = parse_job_a(tmp_path / "x.xlsx")[0] + parse_job_a(tmp_path / "y.xlsx")[0]
    agreement = agreement_job_a(rows)
    assert agreement["shared_profiles"] == 4 and agreement["verdict_agreement"] == 0.75
    assert score_job_a(rows)["all"]["profiles_labelled"] == 4                                   # shared profiles count once


# ---------- job B round trip ----------

def test_job_b_questions_are_parsed_and_checked(tmp_path):
    lookup = {"Papers": (["arXiv id", "title"], [["p1", "T"]])}
    path = build_job_b(tmp_path / "b.xlsx", "Tarun", "TA", ["covered", "one missing", "partly covered"], lookup)
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "READ ME" and wb.sheetnames[1] == "Questions" and "Papers" in wb.sheetnames
    fill(path, "Questions", {
        2: {"question": "How well do models perform on Kannada NLI?", "intent": "result", "complexity": "simple", "language": "Kannada",
            "task": "natural language inference", "expected scope warning": "NO"},
        3: {"question": "XLM-R on XNLI for Kannada?", "intent": "result", "complexity": "simple", "language": "Kannada",
            "expected scope warning": "YES"},                                                                    # no missing condition named
        4: {"question": "Compare mBERT and GPT-4", "intent": "opinion", "complexity": "complex", "expected scope warning": "NO"},
    })
    questions, problems = parse_job_b(path)
    assert [q["intent"] for q in questions] == ["result", "result", "opinion"]
    assert questions[0]["conditions"] == {"task": "natural language inference", "language": "Kannada"} and questions[0]["expect_warning"] is False
    assert any("no missing condition is named" in p for p in problems) and any("opinion" in p for p in problems)
    assert any("planned 'one missing'" in p or "planned" in p for p in problems) is True
    assert B_HEADERS.index("question") == 3


# ---------- job C round trip ----------

def make_pairs():
    pairs = []
    for i, (va, vb, setting_b) in enumerate([(83.6, 79.2, "zero-shot"), (90.0, 80.0, None), (50.0, 55.0, "dev set")]):
        x = prof(2 * i, paper="a", chunk="a:1", model="XLM-R", dataset="XNLI", metric="accuracy", value=va, setting="translate-train")
        y = prof(2 * i + 1, paper="b", chunk="b:1", model="XLM-R", dataset="XNLI", metric="accuracy", value=vb, setting=setting_b)
        pairs.append((f"P0{i + 1}", Pair(x, y, "EXPLAINED", ["setting"], "r", abs(va - vb) / max(va, vb), ("xlmr", "xnli", "accuracy"))))
    return pairs


def test_job_c_round_trip_kappa_gold_and_system_score(tmp_path):
    pairs = make_pairs()
    texts, pages = {"a:1": TABLE, "b:1": TABLE}, {"a:1": 4, "b:1": 9}
    for name in ("one", "two", "judge"):
        build_job_c(tmp_path / f"{name}.xlsx", name, pairs, texts, pages, {"a": "Paper A", "b": "Paper B"})
    wb = load_workbook(tmp_path / "one.xlsx")
    assert wb.sheetnames[:2] == ["READ ME", "Pairs"] and len(C_HEADERS) == 24
    ws = wb["Pairs"]
    assert ws.cell(row=2, column=column_of(ws, "open B")).hyperlink.target == "https://arxiv.org/pdf/b#page=9"
    headers = [c.value for c in ws[1]]
    assert not any("system" in str(h).lower() for h in headers)                                  # the system's verdict is never shown
    rowof = {"P01": 2, "P02": 3, "P03": 4}
    fill(tmp_path / "one.xlsx", "Pairs", {rowof["P01"]: {"verdict": "EXPLAINED", "differs: setting": "Y"},
                                          rowof["P02"]: {"verdict": "GENUINE"}, rowof["P03"]: {"verdict": "EXPLAINED", "differs: setting": "Y"}})
    fill(tmp_path / "two.xlsx", "Pairs", {rowof["P01"]: {"verdict": "EXPLAINED", "differs: setting": "Y", "differs: model size": "Y"},
                                          rowof["P02"]: {"verdict": "NOT COMPARABLE"}, rowof["P03"]: {"verdict": "EXPLAINED", "differs: setting": "Y"}})
    fill(tmp_path / "judge.xlsx", "Pairs", {rowof["P02"]: {"verdict": "GENUINE"}})
    one, p1 = parse_job_c(tmp_path / "one.xlsx")
    two, _ = parse_job_c(tmp_path / "two.xlsx")
    judge, _ = parse_job_c(tmp_path / "judge.xlsx")
    assert p1 == [] and one["P01"]["differs"] == ["setting"] and two["P01"]["differs"] == ["model size", "setting"]
    agreement = agreement_job_c(one, two)
    assert agreement["pairs_labelled_by_both"] == 3 and agreement["disagreements"] == ["P02"] and agreement["percent_agreement"] == pytest.approx(2 / 3, abs=1e-3)
    gold = gold_job_c(one, two, judge)
    assert gold["P02"]["verdict"] == "GENUINE" and gold["P02"]["settled_by_third"]              # the third labeller settled it
    assert gold["P01"]["differs"] == ["setting"]                                                # only what BOTH ticked
    assert "P02" not in gold_job_c(one, two)                                                    # unsettled disagreements are left out
    key = {pid: {"system_verdict": p.verdict, "system_differing": p.differing} for pid, p in pairs}
    score = score_system_c(gold, key)
    assert score["pairs"] == 3 and score["accuracy"] == pytest.approx(2 / 3, abs=1e-3)                    # the system said EXPLAINED for all three
    assert score["per_class"]["GENUINE"]["recall"] == 0.0 and score["condition_attribution"]["exact_match"] == 1.0
