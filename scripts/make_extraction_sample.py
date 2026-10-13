"""A random sample of stored profiles for a field-by-field check (E3, 12 Oct 2026, drawn from the store of that afternoon; E4 is drawn after the last code change).

  python scripts/make_extraction_sample.py build --n 100 --seed 12102026 --dir extraction_test3     # writes items.jsonl + sheet.txt (E4: --dir extraction_test4 --seed ...)
  python scripts/make_extraction_sample.py score [--labels labels_after_fixes.jsonl]   # reads labels.jsonl (the labeller's verdicts) and prints / writes the scores

The sample is uniform over every profile of the 28-paper store that is on none of the earlier hand-checked lists (the two job-A sheets and E1; E2 is a sample of another
store). For each item the labeller sees the paper, the caption, the source row with the table's header, and the repaired fields (the ones the system uses), and writes
for each field OK / WRONG / MISSING (with `imprecise` for a true but vague value: lenient = right, strict = wrong), the same rules as job A. One AI labeller, no kappa.
The sheet is printed in the order of the items; nothing here uses a model.
"""
from __future__ import annotations

import argparse
import json
import math
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

FIELDS = ["model", "dataset", "dataset_version", "metric", "value", "language", "task", "model_size", "setting"]
DIR = ROOT / "data/labelling/private/extraction_test3"


def _earlier_ids() -> set[str]:
    ids: set[str] = set()
    import openpyxl
    for name in ("A_profile_check_Shashank.xlsx", "A_profile_check_Adarsh.xlsx"):
        path = ROOT / "data/labelling" / name
        if path.exists():
            wb = openpyxl.load_workbook(path, read_only=True)
            for r in wb["Profiles"].iter_rows(min_row=2, values_only=True):
                if r[1]:
                    ids.add(r[1])
    for items in sorted((ROOT / "data/labelling/private").glob("extraction_test*/items.jsonl")):
        if items.parent != DIR:
            ids |= {json.loads(x)["profile_id"] for x in open(items, encoding="utf-8") if x.strip() and "profile_id" in x}
    return ids


def _source_row(text: str, value: float, row_label: str | None) -> str:
    """The table line that holds the number as a whole cell number (87.0 is not "87.8"), the one whose label is the grounded row's when several do."""
    hits = []
    for ln in text.split("\n"):
        if not ln.lstrip().startswith("|"):
            continue
        found = False
        for cell in (c.strip() for c in ln.strip().strip("|").split("|")):
            for part in re.split(r"\s*/\s*|\s+", cell):
                m = re.fullmatch(r"[\(\[]?([-+]?\d+(?:\.\d+)?)[%)\]†‡*]*", part)
                if m and abs(float(m.group(1)) - value) < 1e-6:
                    found = True
                    break
            if found:
                break
        if found:
            hits.append(ln)
    if row_label:
        key = re.sub(r"\W", "", row_label).lower()
        for ln in hits:
            first = re.sub(r"\W", "", ln.strip().strip("|").split("|")[0]).lower()
            if key and first and (first == key or key in first or first in key):
                return ln
    return hits[0] if hits else ""


def build(n: int, seed: int) -> None:
    from cgrag.config import get_settings
    from cgrag.labelling.sampling import load_chunk_index
    from cgrag.pipeline.grounding import ground
    from cgrag.ingestion.tables import parse_table
    from cgrag.stores.profile_store import ProfileStore

    cfg = get_settings()
    idx = load_chunk_index(cfg.paths.index_dir / "chunks.jsonl")
    seen = _earlier_ids()
    pool = [p for p in ProfileStore(repaired=True).all() if p.profile_id not in seen and p.value is not None]
    rng = random.Random(seed)
    sample = rng.sample(pool, n)
    DIR.mkdir(parents=True, exist_ok=True)
    lines, items = [], []
    for i, p in enumerate(sample, 1):
        text = idx.texts.get(p.chunk_id, "")
        table = parse_table(text)
        g = ground(p, text) if table else None
        row = _source_row(text, p.value, g.row_label if g else None)
        header = next((ln for ln in text.split("\n") if ln.lstrip().startswith("|")), "")
        item = {"id": f"F{i:03d}", "profile_id": p.profile_id, "paper": p.paper_id, "title": idx.titles.get(p.paper_id, ""),
                "caption": (table.caption if table else text.split("\n")[0])[:240], "header": header[:300], "row": row[:300],
                "cell": ({"row_label": g.row_label, "column": g.column, "block": g.block} if g else None),
                "evidence": (p.evidence or "")[:300], "fields": {f: getattr(p, f) for f in FIELDS}}
        items.append(item)
        lines.append(f"{item['id']} {p.paper_id} {p.chunk_id.split(':')[1]} | {item['caption'][:150]}")
        if table:
            lines.append(f"   HEADER {item['header'][:260]}")
            lines.append(f"   ROW    {item['row'][:260]}")
            if g:
                lines.append(f"   CELL   row={g.row_label!r} column={g.column!r} block={g.block!r}")
        else:
            lines.append(f"   TEXT   {item['evidence'][:260]}")
        lines.append("   FIELDS " + "; ".join(f"{f}={getattr(p, f)!r}" for f in FIELDS))
    (DIR / "items.jsonl").write_text("\n".join(json.dumps(x, ensure_ascii=False) for x in items), encoding="utf-8")
    (DIR / "sheet.txt").write_text("\n".join(lines), encoding="utf-8")
    print(f"{len(items)} items of a pool of {len(pool)} (excluded {len(seen)} earlier ids) -> {DIR}")


def _wilson(k: int, n: int) -> tuple[float, float]:
    if n == 0:
        return 0.0, 0.0
    z, p = 1.96, k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    r = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return 100 * (c - r) / d, 100 * (c + r) / d


def score(labels_file: str = "labels.jsonl") -> None:
    labels = [json.loads(x) for x in open(DIR / labels_file, encoding="utf-8") if x.strip()]
    out = {}
    for strict in (False, True):
        tally = {f: [0, 0, 0] for f in FIELDS}
        full = []
        n_items = 0
        for lab in labels:
            if lab.get("not_a_result"):
                continue
            n_items += 1
            bad, miss, vague = set(lab.get("wrong", [])), set(lab.get("missing", [])), set(lab.get("imprecise", []))
            if strict:
                bad |= vague
            for f in FIELDS:
                t = tally[f]
                if f in bad:
                    t[1] += 1
                    t[2] += 1
                elif f in miss:
                    t[2] += 1
                else:
                    t[0] += 1
            full.append(not (bad | miss))
        tp, fp, fn = (sum(t[i] for t in tally.values()) for i in range(3))
        prec = tp / (tp + fp) if tp + fp else 0
        rec = tp / (tp + fn) if tp + fn else 0
        lo, hi = _wilson(sum(full), len(full))
        out["strict" if strict else "lenient"] = {
            "items": n_items, "completely_right": f"{sum(full)}/{len(full)} = {100 * sum(full) / max(1, len(full)):.1f} %", "ci95": [round(lo, 1), round(hi, 1)],
            "micro_f1": round(100 * 2 * prec * rec / (prec + rec), 1) if prec + rec else 0.0,
            "per_field_f1": {f: round(100 * 2 * t[0] / (2 * t[0] + t[1] + t[2]), 1) if t[0] else 0.0 for f, t in tally.items()},
            "wrong_or_missing_by_field": {f: t[2] for f, t in tally.items()}}
    out["not_a_result"] = sum(1 for lab in labels if lab.get("not_a_result"))
    suffix = "" if labels_file == "labels.jsonl" else "." + Path(labels_file).stem
    (ROOT / f"eval/labels/extraction_{DIR.name}{suffix}.json").write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", nargs="?", default="build", choices=["build", "score"])
    ap.add_argument("--n", type=int, default=100)
    ap.add_argument("--seed", type=int, default=12102026)
    ap.add_argument("--labels", default="labels.jsonl", help="score: the labels file in the folder")
    ap.add_argument("--dir", default="extraction_test3", help="folder under data/labelling/private (extraction_test3 = E3, extraction_test4 = E4)")
    a = ap.parse_args()
    DIR = ROOT / "data/labelling/private" / a.dir
    build(a.n, a.seed) if a.cmd == "build" else score(a.labels)
