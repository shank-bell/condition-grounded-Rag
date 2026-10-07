"""Put AI-drafted job-A verdicts into a copy of a labelling workbook and record where they came from.

The AI (Gemini in a browser, with the prompt in data/labelling/gemini_prompt_A.txt) answers one JSON object per row:
    {"n": 12, "verdict": "SOME_WRONG", "wrong": ["language"], "missing": [], "note": "..."}
Paste its answers, batch after batch, into one text file (data/labelling/gemini_A_answers.txt), then:

  python scripts/fill_job_a_from_ai.py --source "Gemini <model>, web app, 3 Oct 2026"

Every run starts from the untouched original workbook, so it can simply be repeated after each new batch. The result goes to
data/labelling/done/<name>_AIDRAFT_DONE.xlsx (the name still matches A_*_DONE.xlsx for scripts/score_labels.py) plus a
<name>_AIDRAFT_DONE.provenance.json. The READ ME sheet says in red that the verdicts are an AI draft.

These labels are NOT the human answer key: a person has to check a random sample of the rows before the numbers are reported.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties

from cgrag.config import ROOT
from cgrag.labelling import score as S
from cgrag.labelling.sheets import A_FIELDS, A_NOTE, A_VERDICTS, a_flag_header

LAB = ROOT / "data" / "labelling"


def _sha(path: Path) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


def parse_answers(text: str) -> tuple[dict[int, dict], list[str]]:
    """({n: answer}, problems). Lines that hold no JSON object (READY, a markdown fence, chatter) are skipped; a repeated n: the later wins."""
    answers: dict[int, dict] = {}
    problems: list[str] = []
    for k, line in enumerate(text.splitlines(), start=1):
        if line.lstrip().startswith("#"):                 # a note about the file itself, copied into the record
            continue
        m = re.search(r"\{.*\}", line)
        if not m:
            continue
        try:
            obj = json.loads(m.group(0))
        except json.JSONDecodeError as e:
            problems.append(f"answers line {k}: not valid JSON ({e.msg})")
            continue
        try:
            n = int(obj["n"])
        except (KeyError, TypeError, ValueError):
            problems.append(f"answers line {k}: no usable 'n'")
            continue
        verdict = str(obj.get("verdict", "")).strip().upper().replace(" ", "_")
        if verdict not in A_VERDICTS:
            problems.append(f"n={n}: unknown verdict {obj.get('verdict')!r} - row skipped")
            continue
        wrong = [str(x).strip() for x in (obj.get("wrong") or [])]
        missing = [str(x).strip() for x in (obj.get("missing") or [])]
        unknown = [f for f in wrong + missing if f not in A_FIELDS]
        if unknown:
            problems.append(f"n={n}: unknown field name(s) {unknown} ignored")
            wrong, missing = [f for f in wrong if f in A_FIELDS], [f for f in missing if f in A_FIELDS]
        both = set(wrong) & set(missing)
        if both:
            problems.append(f"n={n}: {sorted(both)} listed as wrong AND missing - kept as WRONG")
            missing = [f for f in missing if f not in both]
        if verdict != "SOME_WRONG" and (wrong or missing):
            problems.append(f"n={n}: {verdict} but cells were marked - the marks were dropped")
            wrong, missing = [], []
        if verdict == "SOME_WRONG" and not (wrong or missing):
            problems.append(f"n={n}: SOME_WRONG without a marked cell - please look at this row")
        if n in answers:
            problems.append(f"n={n}: answered twice - the later answer is used")
        imprecise = [f for f in (obj.get("imprecise") or []) if f in A_FIELDS]       # true but too generic to tell two numbers apart
        answers[n] = {"verdict": verdict, "wrong": wrong, "missing": missing, "note": str(obj.get("note") or "").strip(), "imprecise": imprecise}
    return answers, problems


def make_strict(answers: dict[int, dict]) -> dict[int, dict]:
    """The strict reading: a field tagged 'imprecise' counts as WRONG (an ALL_OK row with such a field becomes SOME_WRONG)."""
    out = {}
    for n, a in answers.items():
        a = dict(a)
        if a["imprecise"] and a["verdict"] in ("ALL_OK", "SOME_WRONG"):
            a["wrong"] = a["wrong"] + [f for f in a["imprecise"] if f not in a["wrong"] and f not in a["missing"]]
            a["verdict"] = "SOME_WRONG"
        out[n] = a
    return out


def fill(sheet: Path, answers: dict[int, dict], out: Path, source: str) -> list[str]:
    wb = load_workbook(sheet)
    ws = wb["Profiles"]
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    col_n, col_verdict, col_note = headers["#"], headers["verdict"], headers[A_NOTE]
    col_flag = {f: headers[a_flag_header(f)] for f in A_FIELDS}
    row_of = {int(ws.cell(r, col_n).value): r for r in range(2, ws.max_row + 1) if ws.cell(r, col_n).value is not None}
    problems = []
    for n, a in sorted(answers.items()):
        r = row_of.get(n)
        if r is None:
            problems.append(f"n={n}: the sheet has no such row - skipped")
            continue
        ws.cell(r, col_verdict).value = a["verdict"]
        for f in A_FIELDS:
            ws.cell(r, col_flag[f]).value = "WRONG" if f in a["wrong"] else "MISSING" if f in a["missing"] else None
        ws.cell(r, col_note).value = a["note"] or None
    note = wb["READ ME"]["B3"]
    note.value = (f"AI DRAFT: the verdicts in this file ({len(answers)} of {len(row_of)} rows) were written by {source}, not by a person. "
                  "A human check of a random sample is still to be done.")
    note.font = Font(name="Arial", bold=True, color="C00000", size=10)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return problems


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--sheet", type=Path, default=LAB / "A_profile_check_Shashank.xlsx", help="the untouched original workbook")
    ap.add_argument("--answers", type=Path, default=LAB / "gemini_A_answers.txt")
    ap.add_argument("--prompt", type=Path, default=LAB / "gemini_prompt_A.txt", help="the prompt that was given (its hash is recorded)")
    ap.add_argument("--source", default="Gemini (web app; model version not recorded yet)")
    ap.add_argument("--tag", default="", help="goes into the output name (AIDRAFT-<tag>) so a second AI pass does not overwrite the first")
    ap.add_argument("--strict", action="store_true", help="count a field tagged 'imprecise' as WRONG")
    args = ap.parse_args()

    text = args.answers.read_text(encoding="utf-8")
    answers, problems = parse_answers(text)
    if args.strict:
        answers = make_strict(answers)
    notes = [ln.strip() for ln in text.splitlines() if ln.lstrip().startswith("#")]
    tag = f"-{args.tag}" if args.tag else ""
    out = LAB / "done" / f"{args.sheet.stem}_AIDRAFT{tag}_DONE.xlsx"
    problems += fill(args.sheet, answers, out, args.source + (" [strict: imprecise cells counted as wrong]" if args.strict else ""))

    rows, check = S.parse_job_a(out)          # exactly what scripts/score_labels.py will see
    prov = {"workbook": out.name, "original": args.sheet.name, "labels_by": "AI draft, not a human", "source": args.source,
            "written_at": datetime.now().isoformat(timespec="seconds"), "rows_labelled": len(answers), "rows": sorted(answers),
            "verdicts": dict(Counter(a["verdict"] for a in answers.values())), "strict": args.strict,
            "rows_with_imprecise_cells": sum(bool(a.get("imprecise")) for a in answers.values()),
            "prompt_file": args.prompt.name, "prompt_sha256": _sha(args.prompt), "answers_file": args.answers.name,
            "answers_sha256": _sha(args.answers), "answers_file_notes": notes, "human_check": "pending",
            "problems": problems + check}
    out.with_suffix(".provenance.json").write_text(json.dumps(prov, indent=1, ensure_ascii=False), encoding="utf-8")

    print(f"{len(answers)} rows written to {out}")
    print("verdicts:", prov["verdicts"])
    print(f"scorer reads {len(rows)} labelled rows from it; problems: {len(problems + check)}")
    for p in problems + check:
        print("  PROBLEM", p)


if __name__ == "__main__":
    main()
