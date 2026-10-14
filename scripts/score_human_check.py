"""Score the HUMAN SPOT-CHECK (14 Oct 2026): compare the person's labels with the AI keys and with the system, on the same items.

  python scripts/score_human_check.py                                   # reads data/labelling/human_check/H{1,2,3}_*_DONE.xlsx (any that exist)
  python scripts/score_human_check.py --h1 X.xlsx --h2 Y.xlsx --h3 Z.xlsx

Needs data/labelling/private/human_check_map.json (written by scripts/make_human_check.py: neutral ids -> the AI's label of each item).
Writes eval/labels/human_check.json (aggregates only) and data/labelling/private/human_check_disagreements.md (the items to look at together).
"""
from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from cgrag.config import get_settings
from cgrag.labelling.score import cohen_kappa
from cgrag.labelling.sampling import load_chunk_index
from cgrag.pipeline.contradiction import classify
from cgrag.pipeline.grounding import ground
from cgrag.stores.profile_store import ProfileStore

HC = Path("data/labelling/human_check")
PRIVATE = Path("data/labelling/private")
FIELDS = ["model", "dataset", "dataset_version", "metric", "value", "language", "task", "model_size", "setting"]


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    m = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return (round(100 * (c - m) / d, 1), round(100 * (c + m) / d, 1))


def cell(v) -> str:
    return str(v).strip() if v is not None else ""


def read_sheet(path: Path, sheet: str) -> list[list]:
    wb = load_workbook(path, data_only=True)
    return [list(r) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]


def agreement(a: list[str], b: list[str]) -> dict:
    n = len(a)
    same = sum(x == y for x, y in zip(a, b))
    return {"n": n, "agree": same, "agreement_pct": round(100 * same / n, 1) if n else None,
            "kappa": round(cohen_kappa(a, b), 3) if n and len(set(a) | set(b)) > 1 else None, "wilson_95": wilson(same, n)}


def prf(tp: int, fp: int, fn: int) -> dict:
    p = tp / (tp + fp) if tp + fp else None
    r = tp / (tp + fn) if tp + fn else None
    f = 2 * p * r / (p + r) if p and r else None
    return {"tp": tp, "fp": fp, "fn": fn, "precision": round(100 * p, 1) if p is not None else None,
            "recall": round(100 * r, 1) if r is not None else None, "f1": round(100 * f, 1) if f is not None else None}


def kinds(text: str) -> set[str]:
    out = set()
    for part in str(text).replace(";", ",").split(","):
        part = part.strip().lower()
        if part:
            out.add(part.split("=")[0].strip().replace(" ", "_"))
    return out


def score_h1(path: Path, mapping: dict, notes: list[str]) -> dict:
    rows = read_sheet(path, "Questions")
    human = {}
    for r in rows:
        qid = cell(r[1])
        ans = cell(r[4]).upper()
        if qid and ans in ("YES", "NO"):
            human[qid] = {"warn": ans == "NO", "missing": cell(r[5]), "evidence": cell(r[6]), "sure": cell(r[7]), "note": cell(r[8])}
    results = {}
    for name, fn in (("final13", "eval/labels/questions_gold.final13_main.results.jsonl"), ("frozen", "eval/labels/questions_gold.v2_main.results.jsonl")):
        p = Path(fn)
        if p.exists():
            results[name] = {json.loads(line)["id"]: json.loads(line) for line in p.read_text(encoding="utf-8").splitlines() if line.strip()}
    ids = sorted(human)
    ai = [mapping["questions"][q]["ai_expect_warning"] for q in ids]
    hu = [human[q]["warn"] for q in ids]
    out = {"answered": len(ids), "human_says_warning_needed": sum(hu), "ai_says_warning_needed": sum(ai),
           "human_vs_ai": agreement(["warn" if x else "ok" for x in hu], ["warn" if x else "ok" for x in ai])}
    confusion = Counter((("warn" if h else "ok"), ("warn" if a else "ok")) for h, a in zip(hu, ai))
    out["confusion_human_ai"] = {f"human {h} / ai {a}": n for (h, a), n in confusion.items()}
    for name, res in results.items():
        warned = [bool(res[mapping["questions"][q]["id"]]["warned"]) for q in ids]
        tp = sum(w and h for w, h in zip(warned, hu))
        fp = sum(w and not h for w, h in zip(warned, hu))
        fn = sum((not w) and h for w, h in zip(warned, hu))
        tp_a = sum(w and a for w, a in zip(warned, ai))
        fp_a = sum(w and not a for w, a in zip(warned, ai))
        fn_a = sum((not w) and a for w, a in zip(warned, ai))
        named = [(kinds(human[q]["missing"]), set(k.split("=")[0].strip().lower() for k in res[mapping["questions"][q]["id"]].get("system_missing", []) or []))
                 for q, w, h in zip(ids, warned, hu) if w and h]
        out[f"stage6_{name}_vs_human"] = {**prf(tp, fp, fn), "right_condition_named": f"{sum(bool(a & b) for a, b in named)} of {len(named)}"}
        out[f"stage6_{name}_vs_ai_same_questions"] = prf(tp_a, fp_a, fn_a)
    dis = [q for q, h, a in zip(ids, hu, ai) if h != a]
    for q in dis:
        m = mapping["questions"][q]
        notes.append(f"H1 {q} ({m['id']}): human says {'NO (warn)' if human[q]['warn'] else 'YES (covered)'}, AI key says "
                     f"{'warn' if m['ai_expect_warning'] else 'covered'} (AI missing: {m['ai_missing']}). Question: {m['question']}\n"
                     f"    human missing: {human[q]['missing']!r}; evidence: {human[q]['evidence']!r}; sure: {human[q]['sure']}; note: {human[q]['note']!r}")
    out["low_confidence_answers"] = sum(1 for q in ids if human[q]["sure"].lower() == "low")
    return out


def score_h2(path: Path, mapping: dict, notes: list[str]) -> dict:
    rows = read_sheet(path, "Profiles")
    human = {}
    for r in rows:
        rid = cell(r[1])
        verdict = cell(r[15]).upper()
        if not rid or verdict not in ("ALL_OK", "SOME_WRONG", "NOT_A_RESULT", "CANNOT_CHECK"):
            continue
        flags = {FIELDS[i]: cell(r[16 + i]).upper() for i in range(len(FIELDS)) if cell(r[16 + i])}
        human[rid] = {"verdict": verdict, "flags": flags, "note": cell(r[25]) if len(r) > 25 else ""}
    ids = [r for r in sorted(human) if human[r]["verdict"] in ("ALL_OK", "SOME_WRONG")]

    def human_lenient_right(h) -> bool:
        return h["verdict"] == "ALL_OK" or (h["verdict"] == "SOME_WRONG" and h["flags"] and all(v == "VAGUE" for v in h["flags"].values()))

    def human_strict_right(h) -> bool:
        return h["verdict"] == "ALL_OK"

    ai_labelled = [r for r in ids if mapping["profiles"][r].get("ai_labelled")]

    def ai_flags(r) -> dict:
        m = mapping["profiles"][r]
        d = {f: "WRONG" for f in m["ai_wrong"]}
        d.update({f: "MISSING" for f in m["ai_missing"]})
        d.update({f: "VAGUE" for f in m["ai_imprecise"] if f not in d})
        return d

    def ai_lenient_right(r) -> bool:
        return all(v == "VAGUE" for v in ai_flags(r).values())

    def ai_strict_right(r) -> bool:
        return not ai_flags(r)

    out = {"answered": len(human), "judged_results": len(ids), "not_a_result_by_human": sum(1 for h in human.values() if h["verdict"] == "NOT_A_RESULT"),
           "cannot_check": sum(1 for h in human.values() if h["verdict"] == "CANNOT_CHECK")}
    hl = sum(human_lenient_right(human[r]) for r in ids)
    hs = sum(human_strict_right(human[r]) for r in ids)
    out["human_completely_right_lenient"] = {"k": hl, "n": len(ids), "pct": round(100 * hl / len(ids), 1) if ids else None, "wilson_95": wilson(hl, len(ids))}
    out["human_completely_right_strict"] = {"k": hs, "n": len(ids), "pct": round(100 * hs / len(ids), 1) if ids else None, "wilson_95": wilson(hs, len(ids))}
    if ai_labelled:
        out["lenient_human_vs_ai"] = agreement(["right" if human_lenient_right(human[r]) else "wrong" for r in ai_labelled],
                                               ["right" if ai_lenient_right(r) else "wrong" for r in ai_labelled])
        out["strict_human_vs_ai"] = agreement(["right" if human_strict_right(human[r]) else "wrong" for r in ai_labelled],
                                              ["right" if ai_strict_right(r) else "wrong" for r in ai_labelled])
        ai_l = sum(ai_lenient_right(r) for r in ai_labelled)
        ai_s = sum(ai_strict_right(r) for r in ai_labelled)
        out["ai_completely_right_same_profiles"] = {"lenient": f"{ai_l} of {len(ai_labelled)}", "strict": f"{ai_s} of {len(ai_labelled)}"}
        for reading, bad in (("lenient", {"WRONG", "MISSING"}), ("strict", {"WRONG", "MISSING", "VAGUE"})):
            tp = fp = fn = same = tot = 0
            for r in ai_labelled:
                hf, af = human[r]["flags"], ai_flags(r)
                for f in FIELDS:
                    h_bad, a_bad = hf.get(f) in bad, af.get(f) in bad
                    tot += 1
                    same += h_bad == a_bad
                    tp += h_bad and a_bad
                    fp += h_bad and not a_bad
                    fn += a_bad and not h_bad
            out[f"field_flags_{reading}"] = {"cells": tot, "same_cells_pct": round(100 * same / tot, 1), "human_flagged": tp + fp, "ai_flagged": tp + fn,
                                             "both_flagged": tp, "only_human": fp, "only_ai": fn}
    for r in ai_labelled:
        h = human[r]
        if human_lenient_right(h) != ai_lenient_right(r) or human_strict_right(h) != ai_strict_right(r):
            notes.append(f"H2 {r} (profile {mapping['profiles'][r]['profile_id']}): human {h['verdict']} {h['flags']} note={h['note']!r}; AI wrong/missing/imprecise = "
                         f"{mapping['profiles'][r]['ai_wrong']}/{mapping['profiles'][r]['ai_missing']}/{mapping['profiles'][r]['ai_imprecise']}")
    return out


def norm_verdict(v: str) -> str:
    v = v.strip().upper().replace(" ", "_")
    return v if v in ("EXPLAINED", "GENUINE", "NOT_COMPARABLE") else ""


def system_verdicts(mapping: dict) -> dict:
    cfg = get_settings()
    store = ProfileStore(cfg.paths.profile_db, repaired=True)
    by_id = {p.profile_id: p for p in store.all()}
    chunk_text = load_chunk_index(cfg.paths.index_dir / "chunks.jsonl").texts
    out = {}
    for xid, m in mapping["pairs"].items():
        a, b = by_id[m["profile_a"]], by_id[m["profile_b"]]
        ga, gb = ground(a, chunk_text.get(a.chunk_id, "")), ground(b, chunk_text.get(b.chunk_id, ""))
        verdict, differing, _ = classify(a, b, cfg.contradiction, ga, gb)
        out[xid] = verdict
    return out


def score_h3(path: Path, mapping: dict, notes: list[str]) -> dict:
    rows = read_sheet(path, "Pairs")
    human = {}
    for r in rows:
        xid = cell(r[0])
        v = norm_verdict(cell(r[14]))
        if not xid or not v:
            continue
        differs = [d for d, c in zip(["dataset_version", "model_size", "language", "setting", "other"], r[15:20]) if cell(c).upper() == "Y"]
        human[xid] = {"verdict": v, "differs": differs, "extraction_error": cell(r[21]).upper() == "Y", "confidence": cell(r[22]), "why": cell(r[23])}
    ids = sorted(human)
    ai = [norm_verdict(mapping["pairs"][x]["ai_verdict"]) for x in ids]
    hu = [human[x]["verdict"] for x in ids]
    out = {"answered": len(ids), "human_vs_ai": agreement(hu, ai)}
    conf = Counter(zip(hu, ai))
    out["confusion_human_ai"] = {f"human {h} / ai {a}": n for (h, a), n in sorted(conf.items())}
    for cls in ("GENUINE", "NOT_COMPARABLE", "EXPLAINED"):
        n_ai = sum(a == cls for a in ai)
        out[f"ai_{cls}_also_{cls}_by_human"] = f"{sum(h == cls and a == cls for h, a in zip(hu, ai))} of {n_ai}"
    out["human_not_comparable_flagged_extraction_error"] = f"{sum(human[x]['extraction_error'] for x in ids if human[x]['verdict'] == 'NOT_COMPARABLE')} of " \
                                                           f"{sum(human[x]['verdict'] == 'NOT_COMPARABLE' for x in ids)}"
    out["extraction_error_flagged_by_human"] = sum(human[x]["extraction_error"] for x in ids)
    out["extraction_error_flagged_by_ai"] = sum(bool(mapping["pairs"][x]["ai_extraction_error"]) for x in ids)
    both_expl = [x for x in ids if human[x]["verdict"] == "EXPLAINED" and norm_verdict(mapping["pairs"][x]["ai_verdict"]) == "EXPLAINED"]
    if both_expl:
        canon = {"dataset version / split": "dataset_version", "model size": "model_size"}
        same = sum(set(human[x]["differs"]) == {canon.get(d, d) for d in (mapping["pairs"][x]["ai_differs"] or [])} for x in both_expl)
        out["explained_same_differing_conditions"] = f"{same} of {len(both_expl)}"
    try:
        sysv = system_verdicts(mapping)
        sv = [norm_verdict(sysv[x]) for x in ids]
        out["stage7_vs_human"] = {"accuracy": f"{sum(s == h for s, h in zip(sv, hu))} of {len(ids)}", **agreement(sv, hu)}
        out["stage7_vs_ai_same_pairs"] = {"accuracy": f"{sum(s == a for s, a in zip(sv, ai))} of {len(ids)}"}
        out["always_explained_vs_human"] = f"{sum(h == 'EXPLAINED' for h in hu)} of {len(ids)}"
    except Exception as err:
        out["stage7_vs_human"] = f"not computed: {err}"
    for x in ids:
        a = norm_verdict(mapping["pairs"][x]["ai_verdict"])
        if human[x]["verdict"] != a:
            notes.append(f"H3 {x} ({mapping['pairs'][x]['set']} {mapping['pairs'][x]['pair']}): human {human[x]['verdict']} {human[x]['differs']} "
                         f"(extraction error: {human[x]['extraction_error']}, {human[x]['confidence']}) vs AI {a} {mapping['pairs'][x]['ai_differs']}; human: {human[x]['why']!r}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--h1", type=Path, default=HC / "H1_coverage_questions_DONE.xlsx")
    ap.add_argument("--h2", type=Path, default=HC / "H2_profile_check_DONE.xlsx")
    ap.add_argument("--h3", type=Path, default=HC / "H3_result_pairs_DONE.xlsx")
    ap.add_argument("--out", type=Path, default=Path("eval/labels/human_check.json"))
    args = ap.parse_args()
    mapping = json.loads((PRIVATE / "human_check_map.json").read_text(encoding="utf-8"))
    notes: list[str] = []
    result: dict = {"note": "human spot-check of the AI keys (aggregates only); the person did not see any AI or system label", "seed": mapping.get("seed")}
    for name, path, fn in (("H1_coverage_questions", args.h1, score_h1), ("H2_profile_check", args.h2, score_h2), ("H3_result_pairs", args.h3, score_h3)):
        if path.exists():
            result[name] = fn(path, mapping, notes)
            print(f"\n== {name} ({path.name}) ==")
            print(json.dumps(result[name], indent=1, ensure_ascii=False))
        else:
            print(f"(no file for {name}: {path})")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=1, ensure_ascii=False), encoding="utf-8")
    (PRIVATE / "human_check_disagreements.md").write_text("# Items where the human and the AI key differ\n\n" + "\n\n".join(f"- {n}" for n in notes) + "\n", encoding="utf-8")
    print(f"\nwrote {args.out} and {PRIVATE / 'human_check_disagreements.md'} ({len(notes)} items to look at)")


if __name__ == "__main__":
    main()
